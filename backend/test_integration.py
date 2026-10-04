import os
import sys
import io
import urllib.request
import urllib.parse
import urllib.error
import json
from PIL import Image, ImageDraw
import docx

BASE_URL = "http://127.0.0.1:8000"
DEMO_IMAGE = os.path.join(os.path.dirname(__file__), "..", "demo_samples", "images", "educational_control.png")


def synthetic_document_bytes(format):
    """Convert the clearly labeled fictional fixture without writing personal data."""
    with Image.open(DEMO_IMAGE) as source:
        buf = io.BytesIO()
        source.convert("RGB").save(buf, format=format)
        return buf.getvalue()

passed = 0
failed = 0

import socket
import time
import threading

def is_server_running(host="127.0.0.1", port=8000):
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex((host, port)) == 0

if not is_server_running():
    print("[*] Starting local FastAPI server in background thread for integration testing...")
    sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))
    import main as app_module
    import uvicorn
    server_thread = threading.Thread(
        target=lambda: uvicorn.run(app_module.app, host="127.0.0.1", port=8000, log_level="warning"),
        daemon=True
    )
    server_thread.start()
    for _ in range(40):
        if is_server_running():
            print("[*] Local FastAPI server is UP and responding.")
            break
        time.sleep(0.2)

def log_test(name, success, detail=""):
    global passed, failed
    if success:
        passed += 1
        print(f"[PASS] {name}: {detail}")
    else:
        failed += 1
        print(f"[FAIL] {name}: {detail}")

def http_get(path):
    req = urllib.request.Request(BASE_URL + path, method="GET")
    with urllib.request.urlopen(req, timeout=20) as resp:
        return resp.status, resp.read(), resp.headers

def http_head(path):
    req = urllib.request.Request(BASE_URL + path, method="HEAD")
    with urllib.request.urlopen(req, timeout=20) as resp:
        return resp.status, resp.headers

def post_multipart(path, filename, file_bytes, content_type="application/octet-stream"):
    boundary = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
    body = bytearray()
    body.extend(f"--{boundary}\r\n".encode("utf-8"))
    body.extend(f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode("utf-8"))
    body.extend(f"Content-Type: {content_type}\r\n\r\n".encode("utf-8"))
    body.extend(file_bytes)
    body.extend(f"\r\n--{boundary}--\r\n".encode("utf-8"))

    req = urllib.request.Request(
        BASE_URL + path,
        data=bytes(body),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST"
    )
    try:
        with urllib.request.urlopen(req, timeout=40) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8")

def post_json(path):
    req = urllib.request.Request(BASE_URL + path, data=b"", method="POST")
    try:
        with urllib.request.urlopen(req, timeout=40) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8")

def get_json(path):
    req = urllib.request.Request(BASE_URL + path, method="GET")
    try:
        with urllib.request.urlopen(req, timeout=40) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8")

print("\n=================================================================")
print("RUNNING DOCUMENT FORGERY AI ENTERPRISE VERIFICATION SUITE (v2.1)")
print("=================================================================\n")

# -------------------------------------------------------------
# 1. FRONTEND AND HEALTH ENDPOINTS
# -------------------------------------------------------------
try:
    status, body, _ = http_get("/")
    html_str = body.decode("utf-8")
    log_test("GET / (Frontend SPA delivery)", status == 200 and "<div id=\"root\">" in html_str, f"Status {status}")
except Exception as e:
    log_test("GET / (Frontend SPA delivery)", False, str(e))

try:
    status, _ = http_head("/")
    log_test("HEAD /", status == 200, f"Status {status}")
except Exception as e:
    log_test("HEAD /", False, str(e))

try:
    status, body, _ = http_get("/api/health")
    data = json.loads(body.decode("utf-8"))
    has_fmts = len(data.get("supported_formats", [])) >= 8
    log_test("GET /api/health", status == 200 and data.get("status") == "HEALTHY" and has_fmts, f"Formats: {data.get('supported_formats')}")
except Exception as e:
    log_test("GET /api/health", False, str(e))

# -------------------------------------------------------------
# 2. SAMPLES API
# -------------------------------------------------------------
try:
    status, body, _ = http_get("/api/samples")
    data = json.loads(body.decode("utf-8"))
    samples = data.get("samples", [])
    demo_tagged = all(s.get("is_demo") is True for s in samples)
    log_test("GET /api/samples", status == 200 and len(samples) >= 5 and demo_tagged, f"Found {len(samples)} samples, Demo tags: {demo_tagged}")
except Exception as e:
    log_test("GET /api/samples", False, str(e))

# -------------------------------------------------------------
# 3. EXISTING SAMPLES FORENSIC EVALUATION
# -------------------------------------------------------------
try:
    status, data = post_json("/api/analyze-sample/educational_control.png")
    score = data.get("risk_score")
    auth_status = data.get("authenticity_status")
    conf = data.get("confidence_score")
    exp = bool(data.get("explanation"))
    lim = bool(data.get("limitations"))
    log_test("Sample Analysis (educational_control.png)", status == 200 and exp and lim, f"Score: {score}/100, Tier: {auth_status}, Conf: {conf}%")
except Exception as e:
    log_test("Sample Analysis (educational_control.png)", False, str(e))

try:
    encoded = urllib.parse.quote("original_demo.png")
    status, data = post_json(f"/api/analyze-sample/{encoded}")
    score = data.get("risk_score")
    pages = data.get("total_pages")
    ocr_conf = data.get("ocr_confidence")
    log_test("Sample Analysis (original_demo.png)", status == 200 and pages == 1, f"Score: {score}/100, Pages: {pages}, OCR Conf: {ocr_conf}%")
except Exception as e:
    log_test("Sample Analysis (original_demo.png)", False, str(e))

# -------------------------------------------------------------
# 4. MULTI-FORMAT UPLOAD TESTS (7 COMMON FORMATS)
# -------------------------------------------------------------

# (A) JPEG
try:
    jpg_bytes = synthetic_document_bytes("JPEG")
    status, data = post_multipart("/upload", "test_student.jpg", jpg_bytes, "image/jpeg")
    score = data.get("risk_score")
    status_tier = data.get("authenticity_status")
    log_test("Format: JPEG Ingestion", status == 200 and status_tier is not None, f"Score: {score}/100, Tier: {status_tier}")
except Exception as e:
    log_test("Format: JPEG Ingestion", False, str(e))

# (B) PNG (Screenshot)
try:
    png_bytes = synthetic_document_bytes("PNG")
    status, data = post_multipart("/upload", "screenshot.png", png_bytes, "image/png")
    is_shot = data.get("forgery_analysis", {}).get("image_analysis", {}).get("is_screenshot")
    log_test("Format: PNG Ingestion (Screenshot)", status == 200, f"Screenshot Heuristic: {is_shot}, Tier: {data.get('authenticity_status')}")
except Exception as e:
    log_test("Format: PNG Ingestion (Screenshot)", False, str(e))

# (C) WEBP
try:
    im = Image.new("RGB", (600, 400), color=(240, 240, 240))
    d = ImageDraw.Draw(im)
    d.text((40, 40), "TEST WEBP DOCUMENT IDENTIFIER", fill=(0, 0, 0))
    buf = io.BytesIO()
    im.save(buf, format="WEBP")
    webp_bytes = buf.getvalue()
    status, data = post_multipart("/upload", "generated.webp", webp_bytes, "image/webp")
    log_test("Format: WEBP Ingestion", status == 200, f"Format: {data.get('forgery_analysis', {}).get('image_analysis', {}).get('format')}")
except Exception as e:
    log_test("Format: WEBP Ingestion", False, str(e))

# (D) PDF
try:
    pdf_bytes = synthetic_document_bytes("PDF")
    status, data = post_multipart("/upload", "document.pdf", pdf_bytes, "application/pdf")
    log_test("Format: PDF Ingestion", status == 200 and data.get("total_pages") == 1, f"Pages: {data.get('total_pages')}, Score: {data.get('risk_score')}")
except Exception as e:
    log_test("Format: PDF Ingestion", False, str(e))

# (E) TIFF
try:
    im = Image.new("RGB", (800, 600), color=(255, 255, 255))
    d = ImageDraw.Draw(im)
    d.text((50, 50), "TIFF OFFICIAL DOCUMENT CERTIFICATE", fill=(10, 10, 10))
    buf = io.BytesIO()
    im.save(buf, format="TIFF")
    tiff_bytes = buf.getvalue()
    status, data = post_multipart("/upload", "certificate.tiff", tiff_bytes, "image/tiff")
    log_test("Format: TIFF Ingestion", status == 200, f"Risk Score: {data.get('risk_score')}/100, Tier: {data.get('authenticity_status')}")
except Exception as e:
    log_test("Format: TIFF Ingestion", False, str(e))

# (F) BMP
try:
    im = Image.new("RGB", (700, 500), color=(245, 245, 250))
    d = ImageDraw.Draw(im)
    d.text((50, 50), "BMP SCAN INGESTION VERIFICATION", fill=(20, 20, 20))
    buf = io.BytesIO()
    im.save(buf, format="BMP")
    bmp_bytes = buf.getvalue()
    status, data = post_multipart("/upload", "scan.bmp", bmp_bytes, "image/bmp")
    log_test("Format: BMP Ingestion", status == 200, f"Format: {data.get('forgery_analysis', {}).get('image_analysis', {}).get('format')}")
except Exception as e:
    log_test("Format: BMP Ingestion", False, str(e))

# (G) DOCX
try:
    doc = docx.Document()
    doc.core_properties.author = "Forensic Test Engine"
    doc.add_heading("INVOICE & TAX RECEIPT", level=1)
    doc.add_paragraph("Vendor: Global Tech Solutions Ltd.")
    doc.add_paragraph("GSTIN: 07AAAAA0000A1Z5")
    doc.add_paragraph("Total Amount: $1,450.00")
    doc.add_paragraph("Invoice Number: INV-2026-9812")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()
    status, data = post_multipart("/upload", "invoice.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    doc_type = data.get("document_type")
    has_meta = "docx_properties" in data.get("forgery_analysis", {}).get("metadata_forensics", {})
    log_test("Format: DOCX Ingestion", status == 200 and has_meta, f"Document Type: {doc_type}, DOCX Metadata captured: {has_meta}")
except Exception as e:
    log_test("Format: DOCX Ingestion", False, str(e))

# -------------------------------------------------------------
# 5. FORENSIC INTELLIGENCE & 4-TIER CONTRACT CHECKS
# -------------------------------------------------------------
try:
    status, data = post_json("/api/analyze-sample/educational_control.png")
    required_keys = [
        "authenticity_status", "risk_score", "risk_level",
        "confidence_score", "verdict", "explanation",
        "limitations", "indicators", "ocr_confidence"
    ]
    all_present = all(k in data for k in required_keys)
    valid_status = data.get("authenticity_status") in ("AUTHENTIC", "NEEDS_MANUAL_REVIEW", "SUSPICIOUS", "LIKELY_FORGED", "Likely Genuine", "Needs Manual Review", "Likely Forged")
    log_test("Forensic Tier Contract Fields", all_present and valid_status, f"Tier: {data.get('authenticity_status')}, Conf: {data.get('confidence_score')}%")
except Exception as e:
    log_test("Forensic Tier Contract Fields", False, str(e))

# -------------------------------------------------------------
# 5B. PROOFLY MULTIMODAL & V1 REPORT LOOKUP TESTS
# -------------------------------------------------------------
try:
    status, demo_meta = get_json("/api/proofly/demo-scenarios")
    log_test("Proofly: Demo Scenarios Metadata", status == 200 and len(demo_meta.get("scenarios", [])) >= 4, f"Found {len(demo_meta.get('scenarios', []))} scenarios")
except Exception as e:
    log_test("Proofly: Demo Scenarios Metadata", False, str(e))

try:
    status, demo_res = post_json("/api/proofly/analyze-demo/kavita_whatsapp_scam")
    aid = demo_res.get("analysis_id")
    has_proofly = "proofly" in demo_res
    concern = demo_res.get("proofly", {}).get("assessment", {}).get("concern_level")
    log_test("Proofly: Kavita WhatsApp Scam Execution", status == 200 and has_proofly and bool(aid), f"Concern: {concern}, ID: {str(aid)[:8]}...")

    if aid:
        st_get, v1_data = get_json(f"/api/analysis/{aid}")
        expected_v1_keys = ["document", "assessment", "forensics", "sensitive_fields", "claims", "qr_codes", "urls", "identities", "evidence", "safe_actions", "simple_explanation", "technical_details"]
        all_v1_keys = isinstance(v1_data, dict) and all(k in v1_data for k in expected_v1_keys)
        log_test("Proofly: GET /api/analysis/{id} Contract", st_get == 200 and all_v1_keys, f"Level: {v1_data.get('assessment', {}).get('level') if isinstance(v1_data, dict) else v1_data}")

        st_tech, tech_data = get_json(f"/api/analysis/{aid}/technical")
        log_test("Proofly: GET /api/analysis/{id}/technical", st_tech == 200 and isinstance(tech_data, dict), f"Tech keys: {len(tech_data) if isinstance(tech_data, dict) else tech_data}")

        st_rep, rep_data = get_json(f"/api/analysis/{aid}/report-data")
        log_test("Proofly: GET /api/analysis/{id}/report-data", st_rep == 200 and isinstance(rep_data, dict) and all(k in rep_data for k in expected_v1_keys), f"Report pages: {rep_data.get('document', {}).get('pages') if isinstance(rep_data, dict) else rep_data}")

        st_v1, v1_alt = get_json(f"/api/v1/analysis/{aid}")
        log_test("Proofly: GET /api/v1/analysis/{id} Versioned", st_v1 == 200 and isinstance(v1_alt, dict) and v1_alt.get("schema_version") == "1.0", "Schema version: 1.0")
except Exception as e:
    log_test("Proofly: Analysis Lookup Suite", False, str(e))

# -------------------------------------------------------------
# 6. SECURITY & ROBUSTNESS EDGE CASES
# -------------------------------------------------------------

# (A) Reject Spoofed Extension
try:
    fake_bytes = b"Plain text string not matching JPEG binary signature"
    status, res = post_multipart("/upload", "spoof.jpg", fake_bytes, "image/jpeg")
    log_test("Security: Reject Disguised Magic Bytes", status == 400, f"Rejected HTTP {status}")
except Exception as e:
    log_test("Security: Reject Disguised Magic Bytes", False, str(e))

# (B) Reject Unsupported Extension (.exe)
try:
    exe_bytes = b"MZ\x90\x00BinaryExecutable"
    status, res = post_multipart("/upload", "installer.exe", exe_bytes, "application/octet-stream")
    log_test("Security: Reject Disallowed File Extension", status == 400, f"Rejected HTTP {status}")
except Exception as e:
    log_test("Security: Reject Disallowed File Extension", False, str(e))

# (C) Reject Empty File (0 Bytes)
try:
    status, res = post_multipart("/upload", "empty.jpg", b"", "image/jpeg")
    log_test("Security: Reject 0-Byte Payload", status == 400, f"Rejected HTTP {status}")
except Exception as e:
    log_test("Security: Reject 0-Byte Payload", False, str(e))

# (D) Path Traversal Prevention
try:
    status, res = post_json("/api/analyze-sample/..%2F..%2Fwindows%2Fwin.ini")
    log_test("Security: Prevent Path Traversal in Sample API", status in (400, 404), f"HTTP {status}")
except Exception as e:
    log_test("Security: Prevent Path Traversal in Sample API", False, str(e))

# (E) Corrupted Image Recovery (Magic bytes ok, body corrupted)
try:
    corrupted_jpg = b"\xff\xd8\xff\xe0" + b"\x00" * 40 # valid JPEG header but unreadable stream
    status, res = post_multipart("/upload", "corrupt.jpg", corrupted_jpg, "image/jpeg")
    # Should catch error and return clean HTTP 500 without crashing the server
    log_test("Robustness: Graceful Handling of Corrupted Image", status == 500, f"HTTP {status}")
except Exception as e:
    log_test("Robustness: Graceful Handling of Corrupted Image", False, str(e))

print(f"\n=================================================================")
print(f"VERIFICATION SUMMARY: {passed} PASSED, {failed} FAILED")
print(f"=================================================================\n")

if failed > 0:
    sys.exit(1)
